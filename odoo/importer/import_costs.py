#!/usr/bin/env python3
"""Carga a Odoo 18 los costos de Andrea (Costos.xlsx): materias primas, empaques, mano de obra y lista de materiales.

  python3 import_costs.py check  Costos.xlsx             # solo lee y reporta (no toca Odoo)
  python3 import_costs.py load   Costos.xlsx             # simula contra Odoo
  python3 import_costs.py load   Costos.xlsx --apply     # crea componentes, listas de materiales y costos

Conexión por variables de entorno: ODOO_URL, ODOO_DB, ODOO_USER, ODOO_PASSWORD. Es idempotente.
Hojas leídas: BOARDS, BOXES, DOMES y OTHERS (una receta por producto, en columnas). Tablas y Jaleas son cuentas viejas
y no se leen. El Excel de costos no se sube al repositorio (público): va en odoo/data/ (ignorada por git).

Reglas:
- Proveedores: la columna PROVEEDOR se busca por su título (sirve con o sin ella). Cada proveedor es un contacto de empresa y
  cada componente lleva su proveedor con el costo por kg o por unidad. «Gatherings» es costo propio y no se crea.
- Alimentos: cantidad x costo unitario de la hoja (coincide con el costo por cantidad de la hoja).
- Empaque y envío: la hoja a veces cobra la línea distinto del costo unitario (tabla de madera por tamaño, papel film,
  stickers a Q1...). Se respeta el costo por línea de la hoja: cada costo distinto es un componente propio.
- Mano de obra: componente «Mano de obra (Q)» sin inventario; cantidad = quetzales.
- Los componentes en gramos se guardan en kg (costo por kg) y la lista de materiales usa gramos.
- El costo de cada producto es la suma de sus líneas. Si la hoja no suma sus propias líneas (Mini Dome y Big Dome
  omiten Q2 en el COSTO TOTAL) o cuenta una línea con cantidad 0, se avisa y se sigue la suma de las líneas.
- Tras aplicar se lee el costo guardado en Odoo y se compara con la suma de líneas; si difiere, termina con error.
"""
from __future__ import annotations

import argparse
import importlib.util
import os
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[2]
HOJAS = ("BOARDS", "BOXES", "DOMES", "OTHERS")
TOLERANCIA = 0.02          # quetzales de diferencia aceptada contra el COSTO TOTAL de la hoja
PRECISION_COSTO = 4        # decimales de «Product Price» (un gramo de queso cuesta Q0.0760)
MANO_DE_OBRA = ("MO-GT", "Mano de obra (Q)")
PROPIOS = {"gatherings"}   # «proveedor» que es el propio negocio (tarjeta, complemento de envío): no se crea en Odoo


def _catalogo():
    spec = importlib.util.spec_from_file_location(
        "catalogo", RAIZ / "odoo" / "addons" / "gatherings_cotizacion" / "catalogo.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _clave(texto: str) -> str:
    s = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def _slug(texto: str) -> str:
    return _clave(texto).replace(" ", "-").upper()


@dataclass(frozen=True)
class Componente:
    codigo: str
    nombre: str
    tipo: str            # alimento | empaque | mano_obra
    medida: str          # kg | unidad
    costo: float         # por kg o por unidad (costo por gramo x 1000 si medida == kg)
    proveedor: str = ""  # columna PROVEEDOR de la hoja (vacía en hojas viejas)


@dataclass
class Receta:
    nombre: str
    sku: str | None
    lineas: list[tuple[Componente, float, str]] = field(default_factory=list)   # (componente, cantidad, unidad de la línea)
    costo_total: float = 0.0
    precio_hoja: float = 0.0
    avisos: list[str] = field(default_factory=list)

    def costo(self) -> float:
        total = 0.0
        for comp, q, u in self.lineas:
            total += q * comp.costo / (1000 if comp.medida == "kg" and u == "g" else 1)
        return round(total, 4)


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def leer_recetas(ruta: str) -> list[Receta]:
    import openpyxl
    cat = {_clave(p["nombre"]): p["sku"] for p in _catalogo().PRODUCTOS}
    wb = openpyxl.load_workbook(ruta, data_only=True)
    crudas = []   # (receta, [(zona, base, descriptor, cantidad, unidad, costo_linea, proveedor)])
    vistos: dict[str, str] = {}   # proveedor escrito distinto («Tienda zona 1» / «Tienda Zona 1») = el primero que aparece
    for hoja in HOJAS:
        if hoja not in wb.sheetnames:
            continue
        ws = wb[hoja]
        cab: dict[str, int] = {}
        for celda in ws[2]:      # las columnas se buscan por su título: Andrea agregó PROVEEDOR en la columna B
            if celda.value:
                cab.setdefault(str(celda.value).strip().upper(), celda.column)
        col_unidad, col_prov = cab.get("UNIDAD DE MEDIDA", 4), cab.get("PROVEEDOR")
        costo_pack = cab.get("COSTO POR EMPAQUE")
        if not col_prov and costo_pack and costo_pack > 2 and ws.cell(2, costo_pack - 1).value is None:
            col_prov = costo_pack - 1     # DOMES y OTHERS traen los proveedores pero sin el título PROVEEDOR
        for celda in ws[1]:
            if not celda.value:
                continue
            col, titulo = celda.column, str(celda.value)
            nombre = re.sub(r"\s*\(.*\)\s*$", "", titulo).strip().title().replace("Goat", "GOAT")
            r = Receta(nombre, cat.get(_clave(nombre)))
            filas, zona = [], "alimento"
            for fila in range(4, ws.max_row + 1):
                a = ws.cell(fila, 1).value
                if a is None:
                    continue
                a = str(a).strip()
                if a.startswith("COSTO DE ALIMENTOS"):
                    zona = "empaque"
                    continue
                if a.startswith("COSTO DE EMPAQUES"):
                    zona = "fin"     # lo que sigue son totales de la hoja, no líneas de la receta
                    continue
                if a.startswith(("EMPAQUES", "ALIMENTOS")):
                    continue
                if a == "MANO DE OBRA":
                    r.lineas.append((Componente(MANO_DE_OBRA[0], MANO_DE_OBRA[1], "mano_obra", "unidad", 1.0),
                                     _num(ws.cell(fila, col + 2).value) or 0.0, "unidad"))
                elif a == "COSTO TOTAL":
                    r.costo_total = _num(ws.cell(fila, col + 2).value) or 0.0
                elif a == "PRECIO FINAL":
                    r.precio_hoja = _num(ws.cell(fila, col + 2).value) or 0.0
                elif zona != "fin":
                    q, h = _num(ws.cell(fila, col).value), _num(ws.cell(fila, col + 2).value)
                    unidad_pack = str(ws.cell(fila, col_unidad).value or "").strip().lower()
                    prov = str(ws.cell(fila, col_prov).value or "").strip() if col_prov else ""
                    prov = vistos.setdefault(prov.casefold(), prov) if prov else ""
                    if h and h > 0:
                        desc = str(ws.cell(fila, col + 1).value or "").strip()
                        filas.append((zona, a, desc, q or 0.0, "g" if unidad_pack.startswith("gram") else "unidad", h, prov))
            if filas:
                crudas.append((r, filas))
    return _armar(crudas)


def _sin_cantidad(r: Receta, filas: list) -> list:
    """Líneas con cantidad 0 pero costo: la hoja a veces las suma al COSTO TOTAL (stickers de las cajas) y a veces no.
    Se cuentan como 1 solo si así la suma queda más cerca del COSTO TOTAL de la hoja."""
    ceros = [f for f in filas if not f[3]]
    if not ceros:
        return filas
    con = sum(f[5] for f in filas) + sum(q * c.costo for c, q, _ in r.lineas if c.tipo == "mano_obra")
    sin = con - sum(f[5] for f in ceros)
    usar = abs(con - r.costo_total) <= abs(sin - r.costo_total)
    nombres = ", ".join(f[1] for f in ceros)
    r.avisos.append(f"cantidad 0 con costo en la hoja ({nombres}): " + ("se cuenta como 1" if usar else "no se cuenta"))
    return [(z, b, d, q or 1.0, u, h, p) for z, b, d, q, u, h, p in filas] if usar else [f for f in filas if f[3]]


def _armar(crudas) -> list[Receta]:
    """Una línea = un componente. Un mismo concepto con costos de línea distintos se separa por costo."""
    crudas = [(r, _sin_cantidad(r, filas)) for r, filas in crudas]
    variantes: dict[tuple[str, str], dict[float, set[str]]] = {}
    for _, filas in crudas:
        for zona, base, desc, q, u, h, _ in filas:
            variantes.setdefault((zona, base), {}).setdefault(round(h / q, 6), set()).add(desc)
    proveedor: dict[tuple[str, str], str] = {}
    for r, filas in crudas:
        for zona, base, _, _, _, _, p in filas:
            if p and proveedor.setdefault((zona, base), p) != p:
                r.avisos.append(f"{base}: proveedor distinto en la hoja ({p} y {proveedor[(zona, base)]}); se usa el primero")
    comp: dict[tuple[str, str, float], Componente] = {}
    for (zona, base), costos in variantes.items():
        for costo, descs in costos.items():
            sufijo = ""
            if len(costos) > 1:
                distintos = all(len(d) == 1 and next(iter(d)) not in ("", "aproximación") for d in costos.values()) \
                    and len({next(iter(d)) for d in costos.values()}) == len(costos)
                sufijo = f" ({next(iter(descs))})" if distintos else f" (Q{costo:g} c/u)"
            kg = any(f[4] == "g" for _, fs in crudas for f in fs if f[1] == base)
            prefijo = "MP" if zona == "alimento" else "EMP"
            nombre = base + sufijo
            comp[(zona, base, costo)] = Componente(
                f"{prefijo}-{_slug(nombre)}", nombre, zona, "kg" if kg else "unidad", costo * 1000 if kg else costo,
                proveedor.get((zona, base), ""))
    recetas = []
    for r, filas in crudas:
        for zona, base, desc, q, u, h, _ in filas:
            r.lineas.append((comp[(zona, base, round(h / q, 6))], q, u))
        if abs(r.costo() - r.costo_total) > TOLERANCIA:
            r.avisos.append(f"la suma de líneas Q{r.costo():.2f} difiere del COSTO TOTAL de la hoja Q{r.costo_total:.2f}")
        recetas.append(r)
    return recetas


def resumen(recetas: list[Receta]) -> str:
    comps = {c.codigo for r in recetas for c, _, _ in r.lineas}
    out = [f"{len(recetas)} recetas, {len(comps)} componentes"]
    por_prov: dict[str, set[str]] = {}
    for r in recetas:
        for c, _, _ in r.lineas:
            if c.proveedor:
                por_prov.setdefault(c.proveedor, set()).add(c.codigo)
    if por_prov:
        out.append("  proveedores: " + ", ".join(f"{p} ({len(v)})" for p, v in sorted(por_prov.items(), key=lambda x: -len(x[1]))))
    for r in recetas:
        out.append(f"  {r.nombre:22} {r.sku or '(sin SKU)':18} costo Q{r.costo():8.2f}  hoja Q{r.costo_total:8.2f}  "
                   f"precio hoja Q{r.precio_hoja:g}" + "".join(f"  ⚠ {a}" for a in r.avisos))
    return "\n".join(out)


class Cargador:
    def __init__(self, rpc, aplicar: bool):
        self.rpc, self.aplicar = rpc, aplicar
        self.cat = _catalogo()
        self.uom: dict[str, int] = {}
        self.categorias: dict[str, int] = {}
        self.faltan: set[str] = set()   # productos del catálogo que no existen (para no contarlos dos veces al simular)
        self.contador = {"componentes": 0, "actualizados": 0, "productos": 0, "listas": 0, "costos": 0,
                         "proveedores": 0, "vinculos": 0}
        self.socios: dict[str, int | None] = {}

    def _uoms(self):
        datos = self.rpc.call("ir.model.data", "search_read",
                              [("module", "=", "uom"), ("name", "in", ["product_uom_kgm", "product_uom_gram", "product_uom_unit"])],
                              fields=["name", "res_id"])
        mapa = {d["name"]: d["res_id"] for d in datos}
        self.uom = {"kg": mapa["product_uom_kgm"], "g": mapa["product_uom_gram"], "unidad": mapa["product_uom_unit"]}

    def _precision(self):
        ids = self.rpc.search("decimal.precision", [("name", "=", "Product Price")], limit=1)
        if ids:
            actual = self.rpc.call("decimal.precision", "read", ids, ["digits"])[0]["digits"]
            if actual < PRECISION_COSTO and self.aplicar:
                self.rpc.write("decimal.precision", ids, {"digits": PRECISION_COSTO})

    def _categoria(self, nombre: str) -> int | None:
        if nombre not in self.categorias:
            ids = self.rpc.search("product.category", [("name", "=", nombre)], limit=1)
            self.categorias[nombre] = ids[0] if ids else (
                self.rpc.create("product.category", {"name": nombre}) if self.aplicar else None)
        return self.categorias[nombre]

    def componente(self, c: Componente) -> int | None:
        etiqueta = {"alimento": "Materias primas", "empaque": "Empaque y envío", "mano_obra": "Mano de obra"}[c.tipo]
        vals = {"name": c.nombre, "default_code": c.codigo, "type": "consu", "sale_ok": False, "purchase_ok": True,
                "standard_price": round(c.costo, PRECISION_COSTO), "uom_id": self.uom[c.medida],
                "uom_po_id": self.uom[c.medida], "categ_id": self._categoria(etiqueta),
                "is_storable": c.tipo != "mano_obra"}
        ids = self.rpc.search("product.template", [("default_code", "=", c.codigo)], limit=1)
        if ids:
            self.contador["actualizados"] += 1
            if self.aplicar:
                self.rpc.write("product.template", ids, {k: v for k, v in vals.items() if k not in ("uom_id", "uom_po_id")})
            tmpl = ids[0]
        else:
            self.contador["componentes"] += 1
            if not self.aplicar:
                return None
            tmpl = self.rpc.create("product.template", vals)
        return self.rpc.search("product.product", [("product_tmpl_id", "=", tmpl)], limit=1)[0]

    def _proveedor(self, nombre: str) -> int | None:
        if nombre not in self.socios:
            ids = self.rpc.search("res.partner", [("name", "=", nombre), ("is_company", "=", True)], limit=1)
            if ids:
                self.socios[nombre] = ids[0]
            else:
                self.contador["proveedores"] += 1
                self.socios[nombre] = self.rpc.create(
                    "res.partner", {"name": nombre, "is_company": True, "supplier_rank": 1}) if self.aplicar else None
        return self.socios[nombre]

    def _vincular(self, recetas: list[Receta]):
        """Proveedor de cada componente (product.supplierinfo) con el mismo costo por kg o unidad que el componente."""
        hechos = set()
        for r in recetas:
            for c, _, _ in r.lineas:
                if not c.proveedor or c.codigo in hechos or c.proveedor.casefold() in PROPIOS:
                    continue
                hechos.add(c.codigo)
                socio = self._proveedor(c.proveedor)
                tmpl = self.rpc.search("product.template", [("default_code", "=", c.codigo)], limit=1)
                precio = round(c.costo, PRECISION_COSTO)
                ids = self.rpc.search("product.supplierinfo", [("partner_id", "=", socio), ("product_tmpl_id", "=", tmpl[0])],
                                      limit=1) if socio and tmpl else []
                if ids and abs(self.rpc.call("product.supplierinfo", "read", ids, ["price"])[0]["price"] - precio) < 1e-9:
                    continue
                self.contador["vinculos"] += 1
                if not self.aplicar or not (socio and tmpl):
                    continue
                if ids:
                    self.rpc.write("product.supplierinfo", ids, {"price": precio})
                else:
                    self.rpc.create("product.supplierinfo", {"partner_id": socio, "product_tmpl_id": tmpl[0],
                                                              "price": precio, "min_qty": 0})

    def producto_final(self, sku: str) -> int | None:
        """Plantilla de un producto del catálogo oficial; lo crea si falta (por ejemplo The GOAT o The Break: el addon
        solo crea los productos al instalarse, y el sitio no puede cotizar un producto que Odoo no tiene)."""
        oficial = next(p for p in self.cat.PRODUCTOS if p["sku"] == sku)
        ids = self.rpc.search("product.template", [("default_code", "=", sku)], limit=1)
        if ids:
            return ids[0]
        if sku in self.faltan:
            return None
        self.faltan.add(sku)
        self.contador["productos"] += 1
        if not self.aplicar:
            return None
        campos = self.rpc.call("product.template", "fields_get", [], attributes=["type"])
        vals = {"name": oficial["nombre"], "default_code": sku, "list_price": oficial["precio"], "type": "consu",
                "sale_ok": True, "purchase_ok": False, "is_storable": True}
        vals["taxes_id"] = [(6, 0, [])]   # pequeño contribuyente: sin impuesto en las líneas
        if "gt_web_id" in campos:
            vals.update(gt_web_id=oficial["web_id"], gt_min_qty=oficial["minimo"], gt_unit_label=oficial.get("unidad", ""),
                        gt_quote_mode="review" if oficial.get("revision") else "fixed")
        return self.rpc.create("product.template", vals)

    def correr(self, recetas: list[Receta]) -> dict:
        self._uoms()
        self._precision()
        comp_ids: dict[str, int | None] = {}
        for r in recetas:
            for c, _, _ in r.lineas:
                if c.codigo not in comp_ids:
                    comp_ids[c.codigo] = self.componente(c)
        self._vincular(recetas)
        for p in self.cat.PRODUCTOS:     # todo el catálogo del sitio debe existir en Odoo, tenga o no receta
            self.producto_final(p["sku"])
        for r in recetas:
            if not r.sku:
                print(f"  ⚠ {r.nombre}: sin SKU en el catálogo, se omite")
                continue
            tmpl = self.producto_final(r.sku)
            if tmpl is None:        # simulación de un producto aún no creado
                self.contador["listas"] += 1
                continue
            lineas = [(0, 0, {"product_id": comp_ids[c.codigo], "product_qty": q,
                              "product_uom_id": self.uom[u if c.medida == "kg" else "unidad"]})
                      for c, q, u in r.lineas if comp_ids.get(c.codigo)]
            bom = self.rpc.search("mrp.bom", [("product_tmpl_id", "=", tmpl), ("code", "=", "GT-COSTOS")], limit=1)
            self.contador["listas"] += 1
            if self.aplicar:
                vals = {"product_tmpl_id": tmpl, "product_qty": 1, "type": "normal", "code": "GT-COSTOS",
                        "bom_line_ids": [(5, 0, 0)] + lineas}
                if bom:
                    self.rpc.write("mrp.bom", bom, vals)
                else:
                    self.rpc.create("mrp.bom", vals)
                self.rpc.write("product.template", [tmpl], {"standard_price": round(r.costo(), PRECISION_COSTO)})
                self.contador["costos"] += 1
        return self.contador

    def verificar(self, recetas: list[Receta]) -> list[str]:
        """Compara el costo guardado en Odoo con la suma de líneas de la receta."""
        fallas = []
        for r in recetas:
            if not r.sku:
                continue
            ids = self.rpc.search("product.template", [("default_code", "=", r.sku)], limit=1)
            if not ids:
                fallas.append(f"{r.nombre}: no existe en Odoo")
                continue
            odoo = self.rpc.call("product.template", "read", ids, ["standard_price"])[0]["standard_price"]
            marca = "ok" if abs(odoo - r.costo()) <= TOLERANCIA else "DIFERENTE"
            print(f"  {r.nombre:22} Odoo Q{odoo:8.2f}  líneas Q{r.costo():8.2f}  hoja Q{r.costo_total:8.2f}  {marca}")
            if marca != "ok":
                fallas.append(f"{r.nombre}: Odoo Q{odoo:.2f} contra líneas Q{r.costo():.2f}")
        return fallas + self._verificar_proveedores(recetas)

    def _verificar_proveedores(self, recetas: list[Receta]) -> list[str]:
        fallas, vistos = [], set()
        for r in recetas:
            for c, _, _ in r.lineas:
                if not c.proveedor or c.codigo in vistos or c.proveedor.casefold() in PROPIOS:
                    continue
                vistos.add(c.codigo)
                socio = self.rpc.search("res.partner", [("name", "=", c.proveedor), ("is_company", "=", True)], limit=1)
                tmpl = self.rpc.search("product.template", [("default_code", "=", c.codigo)], limit=1)
                if not (socio and tmpl and self.rpc.search(
                        "product.supplierinfo", [("partner_id", "=", socio[0]), ("product_tmpl_id", "=", tmpl[0])], limit=1)):
                    fallas.append(f"{c.nombre}: sin proveedor {c.proveedor} en Odoo")
        print(f"  proveedores enlazados: {len(vistos) - len(fallas)} de {len(vistos)} componentes")
        return fallas


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("modo", choices=["check", "load"])
    ap.add_argument("excel")
    ap.add_argument("--apply", action="store_true", help="crear de verdad en Odoo (sin esto, solo simula)")
    a = ap.parse_args(argv)
    recetas = leer_recetas(a.excel)
    print(resumen(recetas))
    if a.modo == "check":
        return 0
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from rpc import OdooRPC
    try:
        rpc = OdooRPC(os.environ["ODOO_URL"], os.environ["ODOO_DB"], os.environ["ODOO_USER"], os.environ["ODOO_PASSWORD"])
    except KeyError as e:
        print(f"Falta la variable de entorno {e}", file=sys.stderr)
        return 2
    cargador = Cargador(rpc, a.apply)
    c = cargador.correr(recetas)
    print(("CREADO" if a.apply else "SIMULACIÓN (usá --apply para crear)") + f": {c}")
    if a.apply:
        fallas = cargador.verificar(recetas)
        if fallas:
            print("FALLÓ la verificación:", *fallas, sep="\n  ", file=sys.stderr)
            return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
